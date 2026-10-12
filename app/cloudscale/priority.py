"""Request classification from data/service_priority.xlsx (sheet "Priority")."""
import logging
from pathlib import Path

from . import config

log = logging.getLogger(__name__)

DEFAULTS = {
    "CreateOrder": {"priority": 1, "critical": True},
    "UpdateItem": {"priority": 2, "critical": True},
    "GetCatalog": {"priority": 3, "critical": False},
}


def _truthy(value):
    if isinstance(value, str):
        return value.strip().lower() in {"true", "1", "yes"}
    return bool(value)


def load(path=None):
    path = Path(path or config.PRIORITY_FILE)
    if not path.exists():
        log.warning("%s not found, using built-in priorities", path)
        return {k.lower(): v for k, v in DEFAULTS.items()}

    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = wb["Priority"]
        rows = sheet.iter_rows(values_only=True)
        header = [str(c).strip().lower() if c is not None else "" for c in next(rows)]
        col = {name: header.index(name) for name in ("request_type", "priority", "critical")}

        table = {}
        for row in rows:
            if row[col["request_type"]] is None:
                continue
            req_type = str(row[col["request_type"]]).strip().lower()
            table[req_type] = {
                "priority": int(row[col["priority"]]),
                "critical": _truthy(row[col["critical"]]),
            }
        log.info("loaded %d request types from %s", len(table), path)
        return table
    finally:
        wb.close()


TABLE = load()


def is_critical(request_type):
    """Unknown request types count as critical: the safe side of the Spot rule."""
    if not request_type:
        return True
    entry = TABLE.get(str(request_type).strip().lower())
    return True if entry is None else entry["critical"]


def priority_of(request_type):
    if not request_type:
        return 1
    entry = TABLE.get(str(request_type).strip().lower())
    return 1 if entry is None else entry["priority"]
